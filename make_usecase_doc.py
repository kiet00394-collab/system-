import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def set_cell_background(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_color)
    tcPr.append(shd)

def set_cell_margins(cell, top=140, bottom=140, left=180, right=180):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_border(cell, **kwargs):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        edge_data = kwargs.get(edge)
        if edge_data:
            tag = 'w:{}'.format(edge)
            element = OxmlElement(tag)
            element.set(qn('w:val'), edge_data.get('val', 'single'))
            element.set(qn('w:sz'), str(edge_data.get('sz', 4)))
            element.set(qn('w:space'), '0')
            element.set(qn('w:color'), edge_data.get('color', 'CCCCCC'))
            tcBorders.append(element)
    tcPr.append(tcBorders)

def add_styled_paragraph(doc, text="", style='Normal', space_before=0, space_after=6, line_spacing=1.2, align=WD_ALIGN_PARAGRAPH.LEFT):
    p = doc.add_paragraph(text, style=style)
    p.alignment = align
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = line_spacing
    return p

def add_heading_1(doc, text):
    p = add_styled_paragraph(doc, space_before=14, space_after=6)
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(14)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)
    return p

def add_heading_2(doc, text):
    p = add_styled_paragraph(doc, space_before=10, space_after=4)
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(12.5)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0x2C, 0x52, 0x82)
    return p

def add_heading_3(doc, text):
    p = add_styled_paragraph(doc, space_before=6, space_after=2)
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(12)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
    return p

def add_body(doc, text, bold=False, italic=False):
    p = add_styled_paragraph(doc, space_before=0, space_after=5, line_spacing=1.2)
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(12)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)
    return p

def add_bullet(doc, text, bold_prefix="", level=0):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.2
    p.paragraph_format.left_indent = Inches(0.25 * (level + 1))
    
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = 'Times New Roman'
        r_pre.font.size = Pt(12)
        r_pre.font.bold = True
        r_pre.font.color.rgb = RGBColor(0x11, 0x18, 0x27)
        
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor(0x37, 0x41, 0x51)
    return p

def build_document():
    doc = docx.Document()

    # Thiết lập lề A4 chuẩn học thuật
    for s in doc.sections:
        s.top_margin = Inches(0.85)
        s.bottom_margin = Inches(0.85)
        s.left_margin = Inches(1.0)
        s.right_margin = Inches(0.9)
        s.page_width = Inches(8.27)
        s.page_height = Inches(11.69)

    # ==================== TRANG BÌA TIÊU ĐỀ ====================
    p_inst = add_styled_paragraph(doc, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
    r_inst = p_inst.add_run("ĐẠI HỌC ĐÀ NẴNG\nVIỆN NGHIÊN CỨU VÀ ĐÀO TẠO VIỆT - ANH (VNUK)")
    r_inst.font.name = 'Times New Roman'
    r_inst.font.size = Pt(11)
    r_inst.font.bold = True
    r_inst.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)

    p_sub = add_styled_paragraph(doc, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=18)
    r_sub = p_sub.add_run("KHOA CÔNG NGHỆ THÔNG TIN & TRÍ TUỆ NHÂN TẠO")
    r_sub.font.name = 'Times New Roman'
    r_sub.font.size = Pt(10.5)
    r_sub.font.bold = False
    r_sub.font.color.rgb = RGBColor(0x4B, 0x55, 0x63)

    p_title = add_styled_paragraph(doc, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=16, space_after=8)
    r_title = p_title.add_run("TÀI LIỆU ĐẶC TẢ VÀ THUYẾT MINH BIỂU ĐỒ USE CASE (PHIÊN BẢN 3.0)\n(HỆ THỐNG QUẢN TRỊ & GHÉP PHÒNG KÝ TÚC XÁ THÔNG MINH)")
    r_title.font.name = 'Times New Roman'
    r_title.font.size = Pt(15.5)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)

    p_desc = add_styled_paragraph(doc, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=20)
    r_desc = p_desc.add_run("Nâng cấp toàn diện theo ý kiến phản biện: Bổ sung Quản lý Phòng, Loại phòng, Phê duyệt Business Approve, Quy trình Thanh toán, Quản lý Doanh thu - Lợi nhuận và 5 Business Rules cốt lõi")
    r_desc.font.name = 'Times New Roman'
    r_desc.font.size = Pt(11)
    r_desc.font.italic = True
    r_desc.font.color.rgb = RGBColor(0x4B, 0x55, 0x63)

    # Hộp thông tin nhóm thực hiện
    tbl_meta = doc.add_table(rows=4, cols=2)
    tbl_meta.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_info = [
        ("Đề tài nghiên cứu:", "Hệ thống Quản trị & Ghép phòng KTX theo Độ tương thích và Tối ưu hóa Doanh thu"),
        ("Nhóm sinh viên thực hiện:", "Nhóm 3 (Đại diện các trường thành viên ĐH Đà Nẵng)"),
        ("Sinh viên bảo vệ:", "Nguyễn Duy Thái Bảo (VNUK), Lê Văn Việt (SPKT), Lê Minh Quân (DUE)"),
        ("Phiên bản cập nhật:", "Version 3.0 (Đáp ứng trọn vẹn 100% góp ý của Giảng viên hướng dẫn)")
    ]
    for i, (k, v) in enumerate(meta_info):
        r = tbl_meta.rows[i]
        c0, c1 = r.cells[0], r.cells[1]
        c0.width, c1.width = Inches(2.2), Inches(4.2)
        set_cell_background(c0, "F3F4F6")
        set_cell_background(c1, "FFFFFF")
        set_cell_margins(c0, top=80, bottom=80, left=120, right=120)
        set_cell_margins(c1, top=80, bottom=80, left=120, right=120)
        set_cell_border(c0, top={'val': 'single', 'sz': 4, 'color': 'E5E7EB'}, bottom={'val': 'single', 'sz': 4, 'color': 'E5E7EB'})
        set_cell_border(c1, top={'val': 'single', 'sz': 4, 'color': 'E5E7EB'}, bottom={'val': 'single', 'sz': 4, 'color': 'E5E7EB'})
        
        p0 = c0.paragraphs[0]
        p0.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r0 = p0.add_run(k)
        r0.font.name, r0.font.size, r0.font.bold = 'Times New Roman', Pt(10.5), True
        r0.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)
        
        p1 = c1.paragraphs[0]
        p1.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r1 = p1.add_run(v)
        r1.font.name, r1.font.size = 'Times New Roman', Pt(10.5)
        r1.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)

    add_styled_paragraph(doc, space_after=14)

    # ==================== MỤC 1: BẢNG GIẢI TRÌNH PHẢN BIỆN CỦA THẦY ====================
    add_heading_1(doc, "1. BẢNG TIẾP THU VÀ GIẢI TRÌNH CHI TIẾT Ý KIẾN GÓP Ý CỦA THẦY GIÁO")
    add_body(doc, "Sau khi báo cáo phiên bản ban đầu, Giảng viên đã chỉ ra những điểm còn sơ sài do hệ thống mới chỉ tập trung vào thuật toán ghép cặp thuần túy mà thiếu vắng các mắt xích vận hành doanh nghiệp thực tế của một ký túc xá. Nhóm sinh viên đã nghiêm túc tiếp thu và hoàn thiện triệt để trong phiên bản 3.0:")

    tbl_feedback = doc.add_table(rows=6, cols=3)
    tbl_feedback.alignment = WD_TABLE_ALIGNMENT.CENTER
    fb_headers = ["Góp ý của Thầy giáo", "Hiện trạng phiên bản cũ", "Giải pháp nâng cấp hoàn thiện trong v3.0"]
    fb_widths = [Inches(1.8), Inches(2.2), Inches(2.6)]

    for j, h in enumerate(fb_headers):
        cell = tbl_feedback.rows[0].cells[j]
        cell.width = fb_widths[j]
        set_cell_background(cell, "1B365D")
        set_cell_margins(cell, top=90, bottom=90, left=80, right=80)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(h)
        run.font.name, run.font.size, run.font.bold = 'Times New Roman', Pt(10), True
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    fb_rows = [
        ("1. 'Chưa tính tới lợi nhuận'",
         "Chỉ đo điểm thỏa mãn thói quen và số cặp chặn, không có khía cạnh kinh doanh.",
         "Bổ sung phân hệ Quản lý Doanh thu & Lợi nhuận (UC-23, UC-25), tối ưu tỷ lệ lấp đầy phòng trống (BR-05) và báo cáo biên lãi ròng."),
        ("2. 'Thiếu business approve'",
         "Sinh viên nộp khảo sát là chạy ghép thẳng vào phòng, không qua kiểm duyệt.",
         "Thiết lập quy trình Phê duyệt nghiệp vụ (UC-09, BR-03): Hệ thống duyệt hồ sơ trường A, tư cách sinh viên và kỷ luật trước khi phân phòng."),
        ("3. 'Chưa có quản lý phòng, muốn ở loại phòng nào'",
         "Mặc định phòng đồng nhất, không cho sinh viên chọn loại phòng hay tiện ích.",
         "Bổ sung UC-13 (Hệ thống Quản lý danh mục phòng: 2, 4, 6 người, giá thuê) và UC-07 (Sinh viên chọn loại phòng & máy lạnh theo ngân sách)."),
        ("4. 'Chưa tính tới quá trình thanh toán'",
         "Hoàn toàn bỏ qua bước thu tiền phòng, đặt cọc giữ chỗ và ký hợp đồng.",
         "Bổ sung quy trình thanh toán toàn diện (UC-19, UC-20, UC-21, UC-22): Thanh toán kỳ 1 + Cọc 1 triệu trong 72h, tích hợp VNPay và xuất biên lai."),
        ("5. 'Business rules: sv A tới ktx A tự động ghép, ở tối thiểu 1 kì'",
         "Chưa thể hiện được các quy tắc ràng buộc kinh doanh thực tế của KTX liên kết.",
         "Chuẩn hóa 5 Business Rules nền tảng: BR-01 (Xác thực SV trường A) và BR-02 (Cam kết hợp đồng lưu trú tối thiểu 01 học kỳ).")
    ]

    for i, (gop_y, cu, moi) in enumerate(fb_rows):
        row = tbl_feedback.rows[i + 1]
        bg = "F9FAFB" if i % 2 == 1 else "FFFFFF"
        for j, text_val in enumerate([gop_y, cu, moi]):
            cell = row.cells[j]
            cell.width = fb_widths[j]
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=60, bottom=60, left=80, right=80)
            set_cell_border(cell, top={'val': 'single', 'sz': 4, 'color': 'E5E7EB'}, bottom={'val': 'single', 'sz': 4, 'color': 'E5E7EB'})
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run(text_val)
            run.font.name, run.font.size = 'Times New Roman', Pt(9.5)
            if j == 0:
                run.font.bold = True
                run.font.color.rgb = RGBColor(0x99, 0x1B, 0x1B)
            elif j == 2:
                run.font.bold = True
                run.font.color.rgb = RGBColor(0x16, 0x65, 0x34)
            else:
                run.font.color.rgb = RGBColor(0x37, 0x41, 0x51)

    add_styled_paragraph(doc, space_after=14)

    # ==================== MỤC 2: 5 QUY TẮC NGHIỆP VỤ (BUSINESS RULES) ====================
    add_heading_1(doc, "2. NỀN TẢNG 5 QUY TẮC NGHIỆP VỤ CỐT LÕI (BUSINESS RULES - BR)")
    add_body(doc, "Một hệ thống kỹ thuật phần mềm chuẩn mực không thể hoạt động tùy tiện mà phải tuân thủ nghiêm ngặt các quy tắc nghiệp vụ định hình nên logic ứng dụng:")

    add_bullet(doc, "Ký túc xá A được chỉ định liên kết phục vụ sinh viên thuộc Trường Đại học A (hoặc khối thành viên). Sinh viên trường A khi đăng ký đến KTX A sẽ được hệ thống tự động nhận diện qua định dạng MSSV và email giáo dục trường cấp (@edu.vn). Sau khi xác thực đúng đối tượng, sinh viên sẽ được kích hoạt quyền lợi hỗ trợ tự động ghép phòng bằng thuật toán thông minh.", "BR-01 (Quy tắc Đối tượng Tiếp nhận - Đại học Liên kết): ")
    add_bullet(doc, "Để bảo đảm nếp sống văn hóa nội trú, tránh tình trạng xáo trộn phòng ốc liên tục và bảo toàn dòng tiền tài chính theo kế hoạch của KTX, sinh viên bắt buộc phải cam kết thời hạn hợp đồng lưu trú tối thiểu 01 học kỳ (5 tháng) hoặc trọn năm học. Hệ thống kiên quyết từ chối các đề xuất thuê phòng ngắn hạn vài tuần hoặc vài tháng dở dang.", "BR-02 (Quy tắc Thời hạn Hợp đồng Tối thiểu 1 Kỳ): ")
    add_bullet(doc, "Sinh viên sau khi nộp đơn không được vào thẳng khâu ghép phòng mà hồ sơ bắt buộc phải trải qua bước Phê duyệt nghiệp vụ (Business Approve) của Ban quản lý. Cán bộ KTX thẩm định tư cách sinh viên, lịch sử kỷ luật, xác minh diện chính sách ưu tiên và chỉ chuyển trạng thái 'Approved' cho các hồ sơ đủ điều kiện.", "BR-03 (Quy trình Phê duyệt Nghiệp vụ - Business Approval Workflow): ")
    add_bullet(doc, "Trong thời hạn tối đa 72 giờ kể từ khi hệ thống công bố kết quả phân phòng, sinh viên phải hoàn tất thanh toán tiền phòng học kỳ 1 kèm tiền đặt cọc tài sản (1.000.000 VNĐ). Quá hạn 72h chưa thanh toán, hệ thống tự động hủy kết quả giữ chỗ và chuyển nhượng suất cho sinh viên trong danh sách chờ. Nếu sinh viên tự ý hủy hợp đồng trước 1 kỳ, tiền cọc sẽ không được hoàn trả.", "BR-04 (Quy tắc Đặt cọc & Thời hạn Thanh toán 72h): ")
    add_bullet(doc, "Thuật toán ghép phòng kết hợp: Phân nhóm sinh viên theo cùng nguyện vọng loại phòng (2 người, 4 người, 6 người) trước, sau đó áp dụng Gale-Shapley để ghép cặp tương thích thói quen. Hệ thống ưu tiên lấp đầy 100% số giường của từng phòng trước khi mở phòng mới nhằm giảm thiểu lãng phí chi phí điện nước chung và khấu hao tòa nhà, tối đa hóa biên lợi nhuận ròng cho KTX.", "BR-05 (Quy tắc Tối ưu hóa Lợi nhuận & Tỷ lệ Lấp đầy Phòng): ")

    add_styled_paragraph(doc, space_after=14)

    # ==================== MỤC 3: PHÂN TÍCH TÁC NHÂN (ACTORS) ====================
    add_heading_1(doc, "3. PHÂN TÍCH CÁC TÁC NHÂN HỆ THỐNG (SYSTEM ACTORS)")
    add_body(doc, "Hệ thống v3.0 phân định rõ 03 tác nhân (Actors) với ranh giới trách nhiệm mạch lạc:")

    add_heading_2(doc, "3.1 Tác nhân Sinh viên (Student / Resident - Primary Actor)")
    add_bullet(doc, "Sinh viên thuộc trường Đại học A có nhu cầu nội trú tại Ký túc xá A.", "Đối tượng: ")
    add_bullet(doc, "Đăng ký/đăng nhập tài khoản, chọn loại phòng mong muốn (2, 4, 6 người) và tiện ích theo ngân sách, khai báo 8 tiêu chí thói quen sinh hoạt (có hỗ trợ nạp mẫu Persona), nộp đơn cam kết tối thiểu 1 kỳ, xem bảng xếp hạng gợi ý bạn cùng phòng, thanh toán tiền phòng/cọc qua cổng thanh toán và ký hợp đồng lưu trú điện tử.", "Hành vi: ")

    add_heading_2(doc, "3.2 Tác nhân Hệ thống KTX (System - System/Managing Actor)")
    add_bullet(doc, "Đội ngũ cán bộ quản lý KTX chịu trách nhiệm điều hành toàn diện công tác nghiệp vụ nội trú.", "Đối tượng: ")
    add_bullet(doc, "Cấu hình danh mục loại phòng và giá thuê theo kỳ; thực hiện phê duyệt nghiệp vụ (Business Approve) hồ sơ sinh viên; nạp dữ liệu khảo sát CSV hoặc sinh dữ liệu giả lập có Seed; kích hoạt thuật toán ghép phòng Gale-Shapley và tối ưu lấp đầy phòng; đối soát công nợ thanh toán; theo dõi báo cáo doanh thu, chi phí vận hành, tỷ lệ lấp đầy giường và lợi nhuận ròng của KTX.", "Hành vi: ")

    add_heading_2(doc, "3.3 Tác nhân Cổng thanh toán (Payment Gateway - Secondary System Actor)")
    add_bullet(doc, "Hệ thống trung gian thanh toán trực tuyến (VNPay, Ngân hàng liên kết, Ví điện tử).", "Đối tượng: ")
    add_bullet(doc, "Tiếp nhận yêu cầu giao dịch thanh toán tiền phòng và đặt cọc, xử lý trừ tiền tài khoản an toàn, gửi mã phản hồi kết quả giao dịch về máy chủ KTX để tự động kích hoạt trạng thái 'Đã thanh toán'.", "Hành vi: ")

    add_styled_paragraph(doc, space_after=14)

    # ==================== MỤC 4: BẢNG DANH MỤC 25 USE CASE ====================
    add_heading_1(doc, "4. BẢNG DANH MỤC 25 USE CASE VÀ PHÂN HỆ NGHIỆP VỤ CHUẨN UML")
    add_body(doc, "Toàn bộ hệ thống được quy hoạch thành 06 Phân hệ chức năng (Packages) gồm 25 Use Case:")

    tbl_uc = doc.add_table(rows=26, cols=5)
    tbl_uc.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["Mã UC", "Tên Use Case", "Tác nhân chính", "Phân loại UML", "Mục tiêu và Hành vi nghiệp vụ"]
    col_widths = [Inches(0.8), Inches(1.8), Inches(1.2), Inches(1.1), Inches(2.2)]

    hdr_row = tbl_uc.rows[0]
    for j, h in enumerate(headers):
        cell = hdr_row.cells[j]
        cell.width = col_widths[j]
        set_cell_background(cell, "1B365D")
        set_cell_margins(cell, top=90, bottom=90, left=80, right=80)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(h)
        run.font.name, run.font.size, run.font.bold = 'Times New Roman', Pt(10), True
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    uc_data = [
        # Phân hệ 1
        ("UC-01", "Đăng ký tài khoản KTX", "Sinh viên", "Chính", "Tạo tài khoản KTX gắn với Mã SV trường A và email giáo dục."),
        ("UC-02", "Đăng nhập hệ thống", "Sinh viên, Hệ thống", "Chính", "Xác thực danh tính trước khi truy cập cổng thông tin nội trú."),
        ("UC-03", "Xác thực sinh viên Trường A", "Hệ thống", "<<include>>", "Kiểm tra sinh viên thuộc đối tượng hợp lệ theo BR-01."),
        ("UC-04", "Quản lý hồ sơ cá nhân", "Sinh viên", "Chính", "Cập nhật thông tin liên hệ, giấy tờ tùy thân, người liên hệ khẩn cấp."),
        # Phân hệ 2
        ("UC-05", "Khai báo 8 tiêu chí thói quen", "Sinh viên", "Chính", "Nhập thói quen: ngủ, dậy, gọn gàng, tiếng ồn, học tập, hút thuốc... và tiêu chí top."),
        ("UC-06", "Nạp hồ sơ mẫu Persona", "Sinh viên", "<<extend>>", "Tải nhanh hồ sơ mẫu Thái Bảo (VNUK), Văn Việt (SPKT), Minh Quân (DUE)."),
        ("UC-07", "Chọn loại phòng & tiện ích", "Sinh viên", "Chính", "Chọn phòng 2, 4, 6 người; tiện ích máy lạnh/quạt theo khả năng tài chính."),
        # Phân hệ 3
        ("UC-08", "Nộp đơn đăng ký KTX", "Sinh viên", "Chính", "Nộp đơn kèm cam kết thời hạn lưu trú tối thiểu 1 học kỳ (BR-02)."),
        ("UC-09", "Phê duyệt đơn (Business Approve)", "Hệ thống (System)", "Chính", "Hệ thống kiểm duyệt điều kiện trường A, kỷ luật trước khi phân phòng (BR-03)."),
        ("UC-10", "Xem Xếp hạng bạn cùng phòng", "Sinh viên", "Chính", "Hiển thị Top 8 ứng viên hòa hợp nhất được xếp hạng giảm dần."),
        ("UC-11", "Tính điểm tương thích 2 chiều", "Hệ thống", "<<include>>", "Tính toán ma trận M[a][b] và M[b][a] có xét trọng số ưu tiên 1.5x."),
        ("UC-12", "Xem giải thích tiêu chí chi tiết", "Sinh viên", "<<extend>>", "Xem chi tiết tiêu chí hòa hợp cao (>=0.8) và tiêu chí cảnh báo (<=0.4)."),
        # Phân hệ 4
        ("UC-13", "Quản lý danh mục phòng & giá", "Hệ thống (System)", "Chính", "Cấu hình số phòng, tòa nhà, loại phòng 2/4/6 người, đơn giá thuê mỗi kỳ."),
        ("UC-14", "Quản lý dữ liệu sinh viên KTX", "Hệ thống (System)", "Chính", "Nạp dữ liệu khảo sát từ CSV hoặc sinh dữ liệu giả lập có Seed."),
        ("UC-15", "Thực hiện ghép phòng tự động", "Hệ thống (System)", "Chính", "Kích hoạt giải thuật Gale-Shapley kết hợp tối ưu lấp đầy phòng."),
        ("UC-16", "Chạy thuật toán Gale-Shapley", "Hệ thống", "<<include>>", "Thuật toán đề xuất và chấp nhận hoãn lại đảm bảo ghép cặp ổn định."),
        ("UC-17", "Đếm cặp chặn (Blocking Pairs)", "Hệ thống", "<<include>>", "Quét kiểm tra toàn bộ các cặp chéo chứng minh Blocking Pairs = 0."),
        ("UC-18", "Tối ưu hóa tỷ lệ lấp đầy phòng", "Hệ thống", "<<include>>", "Ưu tiên lấp đầy 100% giường trước khi mở phòng mới để tối ưu lợi nhuận (BR-05)."),
        # Phân hệ 5
        ("UC-19", "Thanh toán tiền phòng & đặt cọc", "Sinh viên", "Chính", "Thanh toán tiền phòng kỳ 1 + tiền cọc 1 triệu VNĐ trong 72h (BR-04)."),
        ("UC-20", "Xử lý giao dịch thanh toán", "Cổng thanh toán", "<<include>>", "Kết nối VNPay/Ngân hàng xử lý giao dịch an toàn và phản hồi trạng thái."),
        ("UC-21", "Quản lý công nợ & xuất biên lai", "Hệ thống (System)", "Chính", "Đối soát các khoản thu, xuất biên lai thu tiền điện tử có mã tra cứu."),
        ("UC-22", "Ký & Tải hợp đồng lưu trú KTX", "Sinh viên", "Chính", "Ký xác nhận hợp đồng thuê phòng số có mã QR hợp lệ."),
        # Phân hệ 6
        ("UC-23", "Báo cáo Doanh thu & Lợi nhuận", "Hệ thống (System)", "Chính", "Thống kê doanh thu theo kỳ, chi phí vận hành, lãi ròng và tỷ lệ lấp đầy."),
        ("UC-24", "Báo cáo đối sánh thuật toán", "Hệ thống (System)", "Chính", "Biểu đồ so sánh điểm thỏa mãn Gale-Shapley vs Ghép ngẫu nhiên."),
        ("UC-25", "Xuất báo cáo & danh sách CSV", "Hệ thống (System)", "<<extend>>", "Trích xuất file CSV/Excel danh sách phòng và bảng doanh thu.")
    ]

    for i, row_data in enumerate(uc_data):
        row = tbl_uc.rows[i + 1]
        bg_col = "F9FAFB" if i % 2 == 1 else "FFFFFF"
        if "<<include>>" in row_data[3]:
            bg_col = "F0FDF4"
        elif "<<extend>>" in row_data[3]:
            bg_col = "FAF5FF"

        for j, val in enumerate(row_data):
            cell = row.cells[j]
            cell.width = col_widths[j]
            set_cell_background(cell, bg_col)
            set_cell_margins(cell, top=50, bottom=50, left=80, right=80)
            set_cell_border(cell, top={'val': 'single', 'sz': 4, 'color': 'E5E7EB'}, bottom={'val': 'single', 'sz': 4, 'color': 'E5E7EB'})
            
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if j != 0 and j != 3 else WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(val)
            run.font.name, run.font.size = 'Times New Roman', Pt(9)
            if j == 0:
                run.font.bold = True
                run.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)
            elif j == 3:
                run.font.bold = True
                if "<<include>>" in val:
                    run.font.color.rgb = RGBColor(0x15, 0x80, 0x3D)
                elif "<<extend>>" in val:
                    run.font.color.rgb = RGBColor(0x7E, 0x22, 0xCE)
                else:
                    run.font.color.rgb = RGBColor(0x03, 0x69, 0xA1)
            else:
                run.font.color.rgb = RGBColor(0x37, 0x41, 0x51)

    add_styled_paragraph(doc, space_after=14)

    # ==================== MỤC 5: ĐẶC TẢ CHI TIẾT CÁC USE CASE MỚI ====================
    add_heading_1(doc, "5. ĐẶC TẢ CHI TIẾT CÁC USE CASE BỔ SUNG THEO Ý KIẾN THẦY GIÁO")

    add_heading_2(doc, "5.1 Đặc tả UC-07: Chọn loại phòng & Tiện ích mong muốn")
    add_bullet(doc, "Sinh viên (Người dùng).", "Tác nhân: ")
    add_bullet(doc, "Cho phép sinh viên chủ động chọn không gian sống (phòng đôi, phòng 4 hay phòng 6) và các tiện ích máy lạnh đi kèm tương ứng với khả năng chi trả tài chính.", "Mục tiêu: ")
    add_bullet(doc, "Sinh viên đã đăng nhập và được xác thực là sinh viên trường A.", "Tiền điều kiện: ")
    add_bullet(doc, "Hệ thống tải danh mục các loại phòng KTX: Phòng 2 người (VIP - 2.5tr/tháng/SV), Phòng 4 người (Tiêu chuẩn - 1.4tr/tháng/SV), Phòng 6 người (Tiết kiệm - 800k/tháng/SV). Sinh viên tích chọn nguyện vọng 1 và nguyện vọng 2 kèm dịch vụ xe/giặt. Nguyện vọng này đóng vai trò bộ lọc phân cụm trước khi thuật toán chạy.", "Luồng chính: ")

    add_heading_2(doc, "5.2 Đặc tả UC-09: Phê duyệt đơn đăng ký KTX (Business Approve)")
    add_bullet(doc, "Hệ thống (System) (Cán bộ phê duyệt).", "Tác nhân: ")
    add_bullet(doc, "Kiểm soát chặt chẽ điều kiện nhập nội trú theo các quy định nghiệp vụ BR-01 và BR-02.", "Mục tiêu: ")
    add_bullet(doc, "Sinh viên đã hoàn tất nộp đơn đăng ký kèm cam kết ở tối thiểu 1 kỳ (trạng thái Pending).", "Tiền điều kiện: ")
    add_bullet(doc, "Cán bộ Hệ thống mở màn hình duyệt đơn. Hệ thống tự động gán nhãn xanh cho hồ sơ hợp lệ. Hệ thống kiểm tra hạn ngạch phòng, đối chiếu kỷ luật cũ và nhấn 'Business Approve' chuyển hồ sơ sang trạng thái 'Approved'. Các hồ sơ này mới được nạp vào hàng đợi chạy Gale-Shapley.", "Luồng chính: ")

    add_heading_2(doc, "5.3 Đặc tả UC-19, UC-20, UC-21: Quy trình Thanh toán Tiền phòng & Đặt cọc")
    add_bullet(doc, "Sinh viên, Cổng thanh toán (VNPay), Hệ thống (System).", "Tác nhân: ")
    add_bullet(doc, "Ràng buộc giữ chỗ bằng tiền cọc 1 triệu VNĐ và thu tiền phòng kỳ 1 trong 72h (BR-04).", "Mục tiêu: ")
    add_bullet(doc, "Đã có kết quả ghép phòng chính thức từ thuật toán Gale-Shapley.", "Tiền điều kiện: ")
    add_bullet(doc, "Sinh viên nhận bảng kê chi phí: Tiền phòng kỳ 1 (5 tháng x 1.400.000 = 7.000.000 đ) + Tiền cọc tài sản (1.000.000 đ) = Tổng 8.000.000 đ. Đồng hồ 72h đếm ngược. Sinh viên quét mã QR VNPay. Cổng thanh toán xác thực SUCCESS. Hệ thống cập nhật 'Đã thanh toán', xuất biên lai điện tử và cho phép ký hợp đồng lưu trú số (UC-22).", "Luồng chính: ")
    add_bullet(doc, "Quá 72h sinh viên không nộp tiền, hệ thống tự động hủy kết quả giữ chỗ, nhường suất cho sinh viên trong danh sách chờ (Waitlist).", "Ngoại lệ: ")

    add_heading_2(doc, "5.4 Đặc tả UC-23: Quản lý Doanh thu, Lợi nhuận & Tỷ lệ Lấp đầy KTX")
    add_bullet(doc, "Hệ thống (System) / Ban Giám đốc KTX.", "Tác nhân: ")
    add_bullet(doc, "Phân tích hiệu quả tài chính và tối ưu hóa bài toán kinh doanh cho ký túc xá.", "Mục tiêu: ")
    add_bullet(doc, "Hệ thống tổng hợp Dashboard tài chính: Tổng doanh thu thu về theo kỳ, tổng quỹ tiền cọc đang giữ, chi phí điện nước/bảo trì/vệ sinh ước tính, và Lợi nhuận ròng dự kiến. Thống kê tỷ lệ lấp đầy phòng (Occupancy Rate). Cảnh báo các phòng bị trống giường dở dang để kích hoạt chính sách dồn phòng ghép bổ sung.", "Luồng chính: ")

    add_styled_paragraph(doc, space_after=14)

    # ==================== MỤC 6: BỘ CÂU HỎI VẤN ĐÁP NÂNG CAO ====================
    add_heading_1(doc, "6. BỘ CÂU HỎI VẤN ĐÁP NÂNG CAO KHI BẢO VỆ VỚI THẦY GIÁO")

    qa_list_v3 = [
        ("Câu 1: Thầy hỏi: 'Tại sao lại cần Use Case Chọn loại phòng mà không ghép ngẫu nhiên vào bất kỳ phòng nào?'",
         "Dạ thưa Thầy, trong ký túc xá thực tế, mỗi sinh viên có hoàn cảnh kinh tế và nhu cầu tiện ích rất khác nhau: có bạn có điều kiện tài chính sẵn sàng chi 2.5 triệu/tháng để ở phòng đôi máy lạnh khép kín, nhưng cũng có bạn sinh viên khó khăn chỉ có thể chi trả 800 ngàn/tháng cho phòng 6 người. Nếu hệ thống chỉ ghép dựa trên giờ ngủ và thói quen mà bỏ qua loại phòng thì sẽ dẫn đến việc ghép 2 người cùng thói quen nhưng một người không đủ tiền đóng tiền phòng. Do đó, việc cho sinh viên chọn loại phòng (UC-07) là bộ lọc phân loại đầu vào bắt buộc trước khi chạy thuật toán tương thích."),

        ("Câu 2: Thầy hỏi: 'Bước Business Approve đóng vai trò gì trong mô hình Use Case?'",
         "Dạ thưa Thầy, 'Business Approve' (UC-09) là cánh cổng kiểm soát nghiệp vụ của Ban quản lý theo Quy tắc BR-03. Không một tổ chức ký túc xá nào cho phép sinh viên cứ điền form là tự động vào ở ngay. KTX phải thẩm định: (1) Sinh viên có đúng là sinh viên trường A không theo BR-01; (2) Sinh viên có vi phạm kỷ luật nội trú cũ không; (3) Sinh viên có cam kết ở đủ tối thiểu 1 học kỳ theo BR-02 không. Chỉ khi hồ sơ được Hệ thống phê duyệt 'Business Approve', sinh viên mới được đưa vào danh sách chính thức để chạy thuật toán phân phòng."),

        ("Câu 3: Thầy hỏi: 'Hệ thống tính toán và tối ưu hóa lợi nhuận cho KTX như thế nào?'",
         "Dạ thưa Thầy, hệ thống quản lý lợi nhuận ở 2 khía cạnh: (1) Về mặt báo cáo tài chính (UC-23), hệ thống tính toán Tổng doanh thu từ các phân khúc phòng (phòng 2 người có biên lợi nhuận cao nhất, phòng 4 người tạo doanh thu ổn định nhất) trừ đi chi phí vận hành (điện nước định mức, bảo vệ, vệ sinh) để ra Lợi nhuận ròng; (2) Về mặt giải thuật (UC-18 & BR-05), hệ thống tối ưu hóa tỷ lệ lấp đầy giường (Occupancy Rate). Thuật toán ưu tiên lấp đầy 100% số giường của từng phòng (ví dụ phòng 4 người phải đủ 4 người), tránh tình trạng nhiều phòng dở dang gây tốn chi phí điện nước và dọn dẹp không cần thiết, nhờ đó tối đa hóa lợi nhuận cho KTX."),

        ("Câu 4: Thầy hỏi: 'Quá trình thanh toán và đặt cọc được thiết kế bảo đảm quyền lợi KTX ra sao?'",
         "Dạ thưa Thầy, theo quy định BR-04 và các Use Case UC-19, UC-20, UC-21: Sau khi có kết quả xếp phòng, sinh viên có thời hạn tối đa 72 giờ để hoàn thành thanh toán tiền phòng kỳ 1 kèm tiền đặt cọc tài sản 1.000.000 VNĐ qua cổng thanh toán VNPay. Tiền cọc này ràng buộc trách nhiệm bảo vệ tài sản và cam kết ở tối thiểu 1 kỳ. Nếu sinh viên tự ý hủy trước hạn, KTX không hoàn cọc để bù đắp chi phí phòng trống. Nếu quá 72h chưa nộp tiền, hệ thống tự động hủy giữ chỗ để nhường cơ hội cho sinh viên khác, đảm bảo KTX không bị treo giường ảo và đảm bảo dòng tiền tài chính."),

        ("Câu 5: Thầy hỏi: 'Tại sao lại quy định sinh viên phải ở tối thiểu 01 học kỳ mà không cho ở theo tháng?'",
         "Dạ thưa Thầy, đây là Quy tắc nghiệp vụ nền tảng BR-02. Ký túc xá sinh viên đại học hoạt động theo học chế tín chỉ theo từng học kỳ (5 tháng). Nếu cho phép sinh viên ở theo tháng và tự do chuyển đi giữa kỳ, KTX sẽ đối mặt với rủi ro phòng bị bỏ trống giữa chừng rất khó tìm người lấp đầy, gây thất thoát doanh thu nghiêm trọng và làm xáo trộn nếp sinh hoạt của các bạn cùng phòng còn lại. Cam kết 1 học kỳ là chuẩn mực vận hành của toàn bộ các ký túc xá đại học.")
    ]

    for q_title, q_ans in qa_list_v3:
        p_q = add_styled_paragraph(doc, space_before=8, space_after=2)
        r_q = p_q.add_run(q_title)
        r_q.font.name, r_q.font.size, r_q.font.bold = 'Times New Roman', Pt(12), True
        r_q.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)

        p_a = add_styled_paragraph(doc, space_before=0, space_after=6, line_spacing=1.2)
        r_a = p_a.add_run(q_ans)
        r_a.font.name, r_a.font.size = 'Times New Roman', Pt(11.5)
        r_a.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)

    # ==================== MỤC 7: KẾT LUẬN ====================
    add_heading_1(doc, "7. KẾT LUẬN")
    add_body(doc, "Với việc nâng cấp lên Phiên bản 3.0, mô hình Use Case của Đề tài đã chuyển hóa từ một sơ đồ thuật toán thuần túy thành một Kiến trúc Giải pháp Quản trị Ký túc xá Toàn diện. Sơ đồ không chỉ đáp ứng chuẩn mực học thuật của ngôn ngữ UML quốc tế (phân định chính xác include/extend), mà còn phản ánh chân thực, sắc sảo toàn bộ các mắt xích quản trị thực tế: từ chọn loại phòng, phê duyệt nghiệp vụ, thanh toán cọc giữ chỗ cho đến bài toán tối ưu hóa doanh thu và lợi nhuận cho KTX.")
    add_body(doc, "Tài liệu này là vũ khí học thuật vững chắc nhất để nhóm tự tin giải trình và đạt điểm số tối đa trước Giảng viên hướng dẫn và Hội đồng chấm đồ án.")

    # Lưu tài liệu vào Downloads và Backup vào scratch docs
    out_dl = r"C:\Users\LENOVO\Downloads\Tai_Lieu_Giai_Thich_So_Do_Use_Case_KTX.docx"
    out_docs = r"C:\Users\LENOVO\.gemini\antigravity\scratch\system-\docs\Tai_Lieu_Giai_Thich_So_Do_Use_Case_KTX.docx"

    doc.save(out_dl)
    doc.save(out_docs)
    print(f"Successfully generated docx v3.0 at:\n1. {out_dl}\n2. {out_docs}")

if __name__ == "__main__":
    build_document()
